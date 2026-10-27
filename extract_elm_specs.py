from pathlib import Path
from docx import Document
from openpyxl import load_workbook
folder = Path.home() / "Downloads" / "ELM-Specs"
terms = ("unit head", "deputy registrar", "registrar", "leave administrator", "designation", "approval", "sanction", "substitute", "role", "director", "dean", "permission", "leave type", "entitlement", "overlap")
if not folder.exists():
    raise SystemExit(f"Folder not found: {folder}")
for file in sorted(folder.glob(".docx")) + sorted(folder.glob(".xlsx")):
    print(f"\n{'=' * 60}\nFILE: {file.name}\n{'=' * 60}")
    rows = []
if file.suffix.lower() == ".docx":
    doc = Document(file)
    rows = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            text = " | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells)
            if text.strip():
                rows.append(text)
else:
    workbook = load_workbook(file, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        print(f"\n[SHEET: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            text = " | ".join(str(value).strip() for value in row if value is not None and str(value).strip())
            if text:
                rows.append(text)
    workbook.close()

for index, line in enumerate(rows, 1):
    if any(term in line.lower() for term in terms):
        print(f"{index}: {line[:500]}")
print("\nExtraction complete.")
