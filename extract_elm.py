import glob
import os

from docx import Document
from openpyxl import load_workbook

d = os.path.expanduser("~/Downloads/ELM-Specs")
out = []
for f in sorted(glob.glob(os.path.join(d, "*"))):
    out.append("\n===== " + os.path.basename(f) + " =====")
    if f.endswith(".docx"):
        doc = Document(f)
        for p in doc.paragraphs:
            if p.text.strip():
                out.append(p.text.strip())
        for t in doc.tables:
            for r in t.rows:
                out.append(" | ".join(dict.fromkeys(c.text.strip().replace("\n", " ") for c in r.cells)))
    elif f.endswith(".xlsx"):
        wb = load_workbook(f, data_only=True)
        for ws in wb:
            out.append("--- sheet " + ws.title)
            for row in ws.iter_rows(values_only=True):
                vals = [str(v).replace("\n", " ") for v in row if v is not None]
                if vals:
                    out.append(" | ".join(vals))
open("elm_specs.txt", "w", encoding="utf-8").write("\n".join(out))
print("lines written:", len(out))
